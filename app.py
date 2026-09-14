import json
import hashlib
import re
from datetime import date, datetime
from io import BytesIO
from pathlib import Path

import streamlit as st


ROOT = Path(__file__).parent
PROMPT_PATH = ROOT / "output" / "gemini_prompt.json"
PRICE_LIST = ROOT / "Thira_Fresh_Farm_Pricelist_updated.xlsx"
PRICE_LIST_CACHE_VERSION = "exact-excel-rows-v3"
PRODUCT_ALIASES = {
    "Green Curly Lettuce": ["selada hijau", "selada keriting"],
    "Lollo Rosso Lettuce": ["selada merah"],
    "Red Oakleaf": ["selada oak merah"],
    "Green Oakleaf": ["selada oak hijau"],
    "Butterhead": ["selada butterhead"],
    "Head Lettuce": ["selada kepala"],
    "Romaine": ["selada romaine"],
    "Baby Romaine": ["selada baby romaine"],
    "Baby Bok Choy": ["pakcoy", "pak choi", "pakcoy baby"],
    "Baby Carrots": ["wortel baby", "baby carrot", "carrot"],
    "Baby Buncis": ["buncis baby"],
    "Kale Curly": ["kale", "kale keriting"],
    "Red Cherry Radish": ["red radish", "lobak merah"],
    "Arugula/Rocula": ["arugula", "rocket", "roket"],
    "Thyme": ["thyme"],
    "Coriander": ["ketumbar", "cilantro"],
    "Rosemary": ["rosemary"],
    "Parsley": ["peterseli"],
    "Tarragon": ["estragon"],
    "Common Mint": ["mint", "daun mint"],
    "Green Radish Micro": ["green radish", "radish hijau", "lobak hijau"],
    "Red Radish Micro": ["red radish micro", "radish merah"],
    "Green Mustard Micro": ["green mustard", "sawi hijau"],
    "Red Amaranth Micro": ["red amaranth", "bayam merah"],
    "Red Cabbage Micro": ["red cabbage", "kubis merah"],
    "Snowpea Peashoots": ["snowpea", "pea shoots", "tunas kacang polong"],
    "Tendril Peashoots": ["tendril", "pea tendril"],
    "Coriander Micro": ["micro coriander", "micro ketumbar"],
    "Marigold": ["marigold flower", "marigold f", "bunga marigold"],
    "Marigold Leaf": ["marigold l", "daun marigold"],
    "Dandelion": ["dandelion flower", "dandelion f", "bunga dandelion"],
    "Dandelion Leaf": ["dandelion l", "daun dandelion"],
    "Nasturtium": ["nasturtium flower", "nasturtium f", "bunga nasturtium"],
    "Nasturtium Leaf": ["nasturtium l", "daun nasturtium"],
    "Dianthus": ["dianthus", "anyelir"],
    "Pansy": ["pansy", "bunga pansy"],
    "Viola": ["viola"],
}
EXAMPLE_CHAT = """14 September 2026

Eastman Kitchen:
Selada merah 3 kg
Selada hijau 2 kg
DT 11.00, Jalan Sudirman No. 10, a.n. Rina
Ambil T1

Tavern Kitchen:
Green Radish Micro 2 pack cut 60 gr
Green Radish Micro 1 pack non cut 90 gr
Baby Carrots 3 kg
DT 13.00 ke Jalan Kaliurang a.n. Budi

Garden Bistro:
Marigold F 12 pcs
Marigold L 1 pack isi 60 pcs
Apple Blossoms 2 pack isi 18 pcs
Viola 10 pcs
DT 15.30, Jalan Diponegoro, atas nama Sari

Hotel Merah:
Green Mustard Micro 2 pack cut 5 gr
Green Mustard Micro 1 pack cut 10 gr
Tendril 2 pack non cut 60 gr
Ketumbar 1 pack 100 gr
PO00678124
DT 09.30, Jalan Prawirotaman, A.n. Dimas

15 September 2026

Resto Baru pesan:
Snowpea 2 pack cut 30 gr
Selada merah 1 kg
Alamat Seturan, a.n. Nina"""


def product_aliases(item: str) -> set[str]:
    return {item.casefold(), *(alias.casefold() for alias in PRODUCT_ALIASES.get(item, []))}


def alias_catalog(items) -> dict:
    return {item: sorted(product_aliases(item)) for item in sorted(set(items))}


def measurement(value: str) -> tuple[float, str] | None:
    match = re.search(r"(\d+(?:[.,]\d+)?)\s*(kg|grams?|gr|g|pcs?|pieces?)", str(value or "").lower())
    if not match:
        return None
    unit = match.group(2)
    unit = "gr" if unit in ("g", "gram", "grams") else "pcs" if unit in ("pc", "pcs", "piece", "pieces") else unit
    return float(match.group(1).replace(",", ".")), unit


def read_pricelist(uploaded) -> list[dict]:
    if not uploaded:
        return []
    try:
        import openpyxl
    except ImportError as exc:
        raise RuntimeError("Install dependencies first: pip install -r requirements.txt") from exc
    workbook = openpyxl.load_workbook(BytesIO(uploaded.getvalue()), data_only=True)
    rows = []
    for sheet in workbook.worksheets:
        headers = [str(cell.value or "").strip().lower() for cell in sheet[1]]
        indexes = {header: index for index, header in enumerate(headers)}
        if "product" not in indexes or "price (rp)" not in indexes:
            continue
        for values in sheet.iter_rows(min_row=2, values_only=True):
            item = values[indexes["product"]]
            price = values[indexes["price (rp)"]]
            if not item or price in (None, ""):
                continue
            package_quantity = str(values[indexes.get("quantity", 0)] or "").strip()
            rows.append({
                "item": str(item).strip(),
                "unit": str(values[indexes.get("pricing unit", 0)] or "pack").strip().lower(),
                "price": float(price),
                "category": str(values[indexes.get("category", 0)] or "").strip(),
                "option": str(values[indexes.get("option", 0)] or "").strip(),
                "package_quantity": package_quantity,
            })
    unique = {}
    for row in rows:
        key = (
            row["item"].casefold(),
            row["unit"].casefold(),
            row["option"].casefold(),
            row["package_quantity"].casefold(),
        )
        unique[key] = row
    return list(unique.values())


def calculate_orders(chats: list[dict], prices: list[dict]) -> list[dict]:
    results = []
    for chat in chats:
        for line in chat["text"].lower().splitlines():
            for row in prices:
                item = str(row.get("item", "")).strip()
                unit = str(row.get("unit", "pack")).strip().lower()
                option = str(row.get("option", "")).strip().lower()
                pack_size = str(row.get("package_quantity", "")).strip().lower()
                if not item or not row.get("price"):
                    continue
                aliases = product_aliases(item) | {item.lower().rstrip("s")}
                if item.lower().startswith("baby "):
                    aliases.update({item.lower()[5:], item.lower()[5:].rstrip("s")})
                base_name = item.lower().removesuffix(" leaf")
                if not item.lower().endswith(" leaf") and re.search(rf"\b{re.escape(base_name)}\s+l\b", line):
                    continue
                if not any(re.search(rf"\b{re.escape(alias)}\b", line) for alias in aliases if alias):
                    continue
                if unit == "pack" and not re.search(r"\bpacks?\b", line):
                    continue
                if unit == "pcs" and re.search(r"\bpacks?\b", line):
                    continue
                if option and option not in ("/pcs", "/pack"):
                    if option == "cut" and re.search(r"\bnon\s*cut\b", line):
                        continue
                    if option == "non cut" and not re.search(r"\bnon\s*cut\b", line):
                        continue
                    if option not in ("cut", "non cut") and option not in line:
                        continue
                expected_measurement = measurement(pack_size)
                line_measurements = {measurement(value) for value in re.findall(r"\d+(?:[.,]\d+)?\s*(?:kg|grams?|gr|g|pcs?|pieces?)", line)}
                if expected_measurement and expected_measurement not in line_measurements:
                    continue
                if unit == "kg":
                    match = re.search(r"(\d+(?:[.,]\d+)?)\s*kg\b", line)
                elif unit == "pcs":
                    match = re.search(r"(\d+(?:[.,]\d+)?)\s*(?:pcs?|pieces?)\b", line)
                else:
                    match = re.search(r"(\d+(?:[.,]\d+)?)\s*packs?\b", line)
                quantity = float(match.group(1).replace(",", ".")) if match else 1
                variant = " / ".join(value for value in (option, pack_size) if value)
                results.append({"item": item, "category": row.get("category", ""), "variant": variant, "pack_quantity": quantity, "unit": unit, "unit_price": row["price"], "total": round(quantity * float(row["price"]), 2)})
                break
    return results


def payload_for(chats: list[dict], order_date: str, calculated_orders: list[dict]) -> dict:
    return {"date": order_date, "messages": chats, "calculated_orders": calculated_orders}


def template_manifest(uploaded) -> dict:
    if not uploaded:
        return {"attached_to_streamlit": False}
    try:
        import openpyxl
    except ImportError as exc:
        raise RuntimeError("Install dependencies first: pip install -r requirements.txt") from exc
    workbook = openpyxl.load_workbook(BytesIO(uploaded.getvalue()), data_only=False)
    result = {"attached_to_streamlit": True, "filename": uploaded.name, "sheets": []}
    for sheet in workbook.worksheets:
        result["sheets"].append({
            "name": sheet.title,
            "max_row": sheet.max_row,
            "max_column": sheet.max_column,
            "headers": [sheet.cell(5, column).value for column in range(4, min(sheet.max_column, 18) + 1)],
            "date_rows": [row for row in range(1, sheet.max_row + 1) if any("2026" in str(sheet.cell(row, column).value) for column in range(1, sheet.max_column + 1))],
        })
    return result


def build_prompt(payload: dict, manifest: dict, mode: str) -> str:
    existing_output = mode != "create"
    edit_mode = mode == "edit"
    has_template = manifest.get("attached_to_streamlit", False)
    template_rule = (
        "Use the attached `rekap_pesanan.xlsx` as the current workbook to UPDATE. Preserve every existing "
        "date section, order, item, value, formula, and formatting. Add only the new input."
        if existing_output
        else
        "Use the attached Excel workbook as a STRUCTURE-ONLY template. Create a fresh workbook with the "
        "same sheet names, layout, merged cells, column order, widths, borders, fills, fonts, alignment, "
        "number formats, row heights, and formulas. The example orders, restaurants, PO codes, dates, "
        "customers, prices, and totals are sample data only: do not copy them into the output."
        if has_template
        else "No Excel template was attached. If you generate a workbook, use a simple order table and say that the original template was not provided."
    )
    return f"""Convert the WhatsApp order messages below into the workbook.

Operation mode: {mode}

{template_rule}

Use every message in the batch. Do not invent missing values. Leave unknown
template cells blank. Keep orders separated by date when the template uses date
sections. Interpret `DT` as delivery time/details and `A.n.` or `a.n.` as atas
nama (the recipient/order name).

Pricing variants are significant: match microgreens by product + `Cut` or
`Non Cut` + the exact gram option; match edible flowers and edible leaves by
product + `pcs` or `pack` + the exact pieces-per-pack option. Never substitute
one variant's price for another variant.

Use the app-calculated prices below as the source for totals. Do not recalculate
them differently. If an item is not in the price list, leave its price blank and
mention it as ambiguous.
Use `name_aliases` to interpret Indonesian/common names, but write the canonical
Excel product name in the workbook. `Marigold F` means Marigold flower;
`Marigold L` means Marigold Leaf. Apply the same F/L rule to matching flower
and leaf products.

For this specific workbook, write the data into the existing table as follows:
- Date section headers are in column F and look like `Rabu,1 Juli 2026`.
- The table headers are on the row immediately below each date header.
- The unlabeled column D is the global order sequence: it continues across
  every date and must not reset.
- Column E, headed `NO`, is the per-date order number: reset it to 1 at the
  beginning of each date section and increment it only within that date.
- When one batch contains multiple dates, calculate column D globally but
  calculate column E separately for each date.
- Column F is normally the restaurant, but it can also contain PO codes,
  delivery text, addresses, or `DT... a.n. ...` metadata. Never create a new
  restaurant from a PO/DT/address value.
- Column H is `Nama Item`; column I is item type; J is quantity; K is ready
  stock; L is unit; M is unit price; N is item total; O is order total;
  P is delivery/ongkir detail; Q is pickup code such as T1/T2.
- Use the template's currency formatting for every monetary cell. Store prices
  and totals as numeric Excel values, not text, and display them in the same
  format as the example: `Rp35,000.00` with comma thousands separators and two
  decimal places. Apply it to Harga, Harga Item, Jumlah Penjualan Produk,
  Ongkir, and Harga Beli where those cells contain amounts.
- Column Q (`Ambil`) must remain blank unless the new chat explicitly contains
  an `Ambil` value such as `T1` or `T2`. Never default it to `T1`, `T2`, or any
  other value.
- A non-empty order number in column D starts a new order. Blank order-number
  cells continue the previous order until the next order number.
- Build date sections only for dates found in the new input. Put each chat item
  into its matching new date section. Put delivery time,
  address, PO code, atas nama, and ongkir in the existing row/field used by the
  template; do not move them into the restaurant field.
- {"For an update, keep all existing data. For the same date, append new rows after the last existing order in that date section. For a new date, add a complete new section below the last existing section." if existing_output else "Start with blank data rows. Never copy sample order values from the attached workbook. If several messages contain the same date, combine them in that one new date section; do not overwrite one message with another."}
- {"EDIT MODE: match each chat order to an existing row using date + restaurant + item + variant (Cut/Non Cut and package quantity, or pcs/pack and pieces). If it matches, keep unchanged values unchanged and update only fields explicitly changed in the chat. Do not create a duplicate row. If no exact match exists, do not add it; list it as unmatched in the note." if edit_mode else ""}
- {"When adding a new date section, copy the complete existing section's date row, header row, body-row borders, fills, fonts, alignment, number formats, row heights, and formulas." if existing_output else "For every new date, create a complete section by copying only the template's structure: date row, header row, body-row borders, fills, fonts, alignment, number formats, row heights, and formulas. Then fill the blank body rows with the new input."}
- Never use a blank row without copying its neighboring body-row styles. Every
  new item row must have the same borders and formatting as the template body.
- Blank rows that remain inside the table are still table rows: keep their full
  borders across every table column, even when every cell is empty. Do not
  leave unbordered gaps between orders or at the bottom of a date section.
- Copy the complete body-row style into unused rows before clearing their
  values; blank cells must retain borders, fills, alignment, and number formats.

Do not interpret this as permission to edit a user's local computer. Return one
downloadable workbook. {"This is an edit/update: preserve all unrelated existing data." if edit_mode else "This is an update: the output must contain all previous data plus the new data." if existing_output else "This is initial creation: the output must contain only the new data in the copied template layout."}

Critical file rule:
- Produce exactly one workbook named `rekap_pesanan.xlsx`.
- Do not create files named V1, V2, Final, New, timestamped, or duplicate files.
- Do not add worksheets, columns, helper files, or redesign the workbook.
- {"Do not delete, replace, or recreate the existing 4 July table when adding 5 July. Do not delete any prior date." if existing_output else "Do not carry over sample date tables or sample orders unless that date and order are present in the new input."}
- Return only the completed workbook and a short note about ambiguous values.

Before returning, verify that the fresh workbook contains the input dates,
restaurant names, item names, quantities, prices, DT times, and A.n. names,
while none of the template's sample order values remain.

Structured input:
{json.dumps(payload, ensure_ascii=False, indent=2)}

Template manifest extracted by the prompt builder:
{json.dumps(manifest, ensure_ascii=False, indent=2)}

The prompt builder cannot attach files to Gemini Web. You must attach the actual
Excel file separately in the same Gemini conversation. The manifest above is
only a fallback description and cannot reproduce borders or formatting by itself.
"""


def _date_matches(text: str, selected: str) -> bool:
    match = re.search(r"(\d{1,2})\s+([A-Za-z]+)\s+(\d{4})", text)
    if not match:
        return False
    months = {"januari": 1, "februari": 2, "maret": 3, "april": 4, "mei": 5, "juni": 6, "juli": 7, "agustus": 8, "september": 9, "oktober": 10, "november": 11, "desember": 12}
    try:
        wanted = datetime.strptime(selected, "%Y-%m-%d").date()
        return (int(match.group(3)), months.get(match.group(2).lower()), int(match.group(1))) == (wanted.year, wanted.month, wanted.day)
    except (ValueError, TypeError):
        return selected.lower() in text.lower()


def order_text_from_excel(uploaded, selected_date: str) -> str:
    try:
        import openpyxl
    except ImportError as exc:
        raise RuntimeError("Install dependencies first: pip install -r requirements.txt") from exc

    workbook = openpyxl.load_workbook(BytesIO(uploaded.getvalue()), data_only=True)
    sheet = workbook.active
    date_pattern = re.compile(r"(?:Monday|Tuesday|Wednesday|Thursday|Friday|Saturday|Sunday|Senin|Selasa|Rabu|Kamis|Jumat|Sabtu|Minggu)", re.I)
    active_date = ""
    orders = []
    current = None
    for row in sheet.iter_rows():
        values = [cell.value for cell in row]
        text = " ".join(str(value) for value in values if value not in (None, ""))
        if date_pattern.search(text) and re.search(r"\d{1,2}.*\d{4}", text):
            active_date = text.strip()
            continue
        if not active_date or (selected_date and not _date_matches(active_date, selected_date)):
            continue
        if not text:
            continue
        # In the example, D starts a new order. F contains either the restaurant
        # or metadata such as PO/DT/address/atas-nama for that same order.
        cells = values[5:18]
        restaurant, _, item, _, quantity, _, unit, price, item_amount, sales, delivery, pickup, buy_price = (cells + [None] * 13)[:13]
        if restaurant == "Resto" or item == "Nama Item":
            continue
        if values[3] not in (None, ""):
            current = {"restaurant": "", "items": [], "metadata": [], "total": None}
            orders.append(current)
        if current is None:
            continue
        restaurant_text = str(restaurant or "").strip()
        if restaurant_text:
            if re.match(r"^(PO\w+|DT\s*)", restaurant_text, re.I):
                current["metadata"].append(restaurant_text)
            elif not current["restaurant"]:
                current["restaurant"] = restaurant_text
        if item and str(item).strip().lower() != "ongkir":
            current["items"].append({"item": item, "quantity": quantity, "unit": unit, "price": price, "amount": item_amount})
        if item and str(item).strip().lower() == "ongkir":
            current["metadata"].append(f"ongkir: {price or delivery or pickup}")
        if sales not in (None, ""):
            current["total"] = sales
    if not orders:
        return f"No order rows found for date containing `{selected_date}`."

    lines = [f"ORDER LIST - {selected_date}"]
    for number, order in enumerate(orders, 1):
        details = [f"{number}. {order['restaurant'] or '(restaurant missing)'}"]
        items = "; ".join(f"{item['item']} x {item['quantity'] or '?'} {item['unit'] or ''} ({item['amount'] or item['price'] or 'price missing'})" for item in order["items"])
        if items:
            details.append(f"items: {items}")
        if order["total"]:
            details.append(f"total: {order['total']}")
        if order["metadata"]:
            delivery = " | ".join(order["metadata"])
            details.append(f"details: {delivery}")
            delivery_time = re.search(r"DT\s*([0-9:.]+)", delivery, re.I)
            recipient = re.search(r"(?:A\.n\.?|a\.n\.?)\s*(.+)$", delivery)
            if delivery_time:
                details.append(f"delivery time: {delivery_time.group(1)}")
            if recipient:
                details.append(f"atas nama: {recipient.group(1).strip()}")
        lines.append(" | ".join(details))
    return "\n".join(lines)


st.set_page_config(page_title="Gemini Excel Prompt Builder", layout="wide")
st.title("Gemini Excel Prompt Builder")
st.caption("Batch WhatsApp chats into one strict prompt for Gemini Web.")

with st.sidebar:
    st.header("Workbook inputs")
    operation_label = st.radio("Operation", ["Create new workbook", "Add new orders", "Edit existing orders"])
    mode = {"Create new workbook": "create", "Add new orders": "add", "Edit existing orders": "edit"}[operation_label]
    template = st.file_uploader("Initial example/template (optional)", type=["xlsx"], key="template")
    existing = st.file_uploader("Current rekap_pesanan.xlsx for updates (optional)", type=["xlsx"], key="existing")
    pricelist_upload = st.file_uploader("Pricelist Excel (optional)", type=["xlsx"], key="pricelist")
    st.info("For later entries, upload the latest rekap_pesanan.xlsx here. Attach the downloaded workbook manually in Gemini Web.")
    if mode in ("add", "edit") and not existing:
        st.warning("Upload the current rekap_pesanan.xlsx for this operation mode.")

left, right = st.columns(2)
with left:
    st.subheader("1. Add chats in bulk")
    order_date = st.date_input("Order date", value=date.today())
    chat_files = st.file_uploader("Upload multiple .txt chat files", type=["txt"], accept_multiple_files=True)
    pasted = st.text_area("Or paste multiple chats", height=260, placeholder="Paste all messages here, or separate batches with headings.")
    st.text_area("Recommended chat format (editable/copyable example)", EXAMPLE_CHAT, height=420, key="example_chat")
    chats = [{"name": file.name, "text": file.getvalue().decode("utf-8", errors="replace")} for file in chat_files]
    if pasted.strip():
        chats.append({"name": "pasted-chat", "text": pasted.strip()})
    price_source = pricelist_upload or (BytesIO(PRICE_LIST.read_bytes()) if PRICE_LIST.exists() else None)
    try:
        price_rows = read_pricelist(price_source)
    except (OSError, RuntimeError, ValueError) as exc:
        st.warning(f"Pricelist could not be read automatically: {exc}")
        price_rows = []
    if price_source is not None and not price_rows:
        st.warning("No price rows were detected. Check that the workbook has Product and Price (Rp) columns.")
    source_bytes = pricelist_upload.getvalue() if pricelist_upload else (PRICE_LIST.read_bytes() if PRICE_LIST.exists() else b"")
    source_key = f"{PRICE_LIST_CACHE_VERSION}:{hashlib.sha256(source_bytes).hexdigest()}"
    if st.session_state.get("price_source") != source_key:
        st.session_state.price_source = source_key
        st.session_state.price_rows = price_rows
    edited_prices = st.session_state.get("price_rows", price_rows)
    st.subheader("Price list")
    st.caption("Choose a product and change its price when needed.")
    if edited_prices:
        product_names = sorted({row["item"] for row in edited_prices})
        product_search = st.text_input("Search product", placeholder="Try selada merah, arugula, tendril, or snowpea")
        filtered_products = [name for name in product_names if not product_search.strip() or any(product_search.casefold() in alias for alias in product_aliases(name))]
        if not filtered_products:
            st.warning("No matching products.")
            filtered_products = product_names
        selected_product = st.selectbox("Product", filtered_products, format_func=lambda name: f"{PRODUCT_ALIASES.get(name, [''])[0]} ({name})" if name in PRODUCT_ALIASES else name)
        variant_indices = [index for index, row in enumerate(edited_prices) if row["item"] == selected_product]
        variant_labels = [f"{edited_prices[index].get('option', '')} / {edited_prices[index].get('package_quantity', '')} / {edited_prices[index]['unit']}" for index in variant_indices]
        selected_variant = st.selectbox("Variant", range(len(variant_labels)), format_func=lambda index: variant_labels[index])
        selected_index = variant_indices[selected_variant]
        selected_price = st.number_input("New price (Rp)", min_value=0.0, value=float(edited_prices[selected_index].get("price", 0)), step=500.0)
        if st.button("Update price"):
            edited_prices[selected_index]["price"] = selected_price
            st.session_state.price_rows = edited_prices
            st.success(f"Updated {edited_prices[selected_index]['item']}.")
    else:
        st.warning("No products loaded. Upload the pricelist Excel file.")
    calculated_orders = calculate_orders(chats, edited_prices)
    if calculated_orders:
        st.code(json.dumps(calculated_orders, ensure_ascii=False, indent=2), language="json")
    payload = payload_for(chats, str(order_date), calculated_orders)
    payload["operation"] = mode
    payload["name_aliases"] = alias_catalog(row["item"] for row in edited_prices)
    try:
        source_workbook = existing or template
        manifest = template_manifest(source_workbook)
        prompt = build_prompt(payload, manifest, mode)
    except (OSError, RuntimeError, ValueError) as exc:
        st.error(str(exc))
        manifest = {"attached_to_streamlit": False}
        prompt = build_prompt(payload, manifest, mode)

    st.subheader("2. Generate order list from the Excel")
    summary_date = st.text_input("Date to find in the workbook", value=order_date.strftime("%Y-%m-%d"))
    if template and st.button("Generate dated order list"):
        try:
            st.session_state.order_list = order_text_from_excel(template, summary_date)
        except (OSError, RuntimeError, ValueError) as exc:
            st.error(str(exc))
    if "order_list" in st.session_state:
        st.text_area("Copyable order list", st.session_state.order_list, height=300)
        st.download_button("Download order_list.txt", st.session_state.order_list, "order_list.txt", "text/plain")

with right:
    st.subheader("3. Prompt for Gemini Web")
    if existing:
        st.download_button("Download current rekap for Gemini", existing.getvalue(), existing.name, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    elif template:
        st.download_button("Download Excel template for Gemini", template.getvalue(), template.name, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    st.code(prompt, language="text")
    artifact = {"prompt": prompt, "payload": payload}
    st.download_button("Download gemini_prompt.json", json.dumps(artifact, ensure_ascii=False, indent=2), "gemini_prompt.json", "application/json")
    st.download_button("Download order_payload.json", json.dumps(payload, ensure_ascii=False, indent=2), "order_payload.json", "application/json")
    if st.button("Save prompt files"):
        PROMPT_PATH.parent.mkdir(exist_ok=True)
        PROMPT_PATH.write_text(json.dumps(artifact, ensure_ascii=False, indent=2), encoding="utf-8")
        st.success("Saved output/gemini_prompt.json")
